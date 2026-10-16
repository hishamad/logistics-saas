import io
from concurrent.futures import ThreadPoolExecutor

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from freight.admin import main
from freight.api.app import create_app
from freight.fixtures import demo_snapshot
from freight.onboarding import saudi_network
from freight.sheets import workbook

WRITE = {"X-Requested-With": "freight"}


def template(kind: str) -> bytes:
    return workbook(demo_snapshot(), kind, template=True)  # type: ignore[arg-type]


def test_the_saudi_network_is_complete():
    sites, lanes = saudi_network()
    assert len(sites) >= 20
    assert len(lanes) == len(sites) * (len(sites) - 1)  # every directed pair has a road time


def test_a_carrier_is_created_from_its_own_files_and_can_sign_in(tmp_path, monkeypatch):
    (tmp_path / "fleet.xlsx").write_bytes(template("fleet"))
    (tmp_path / "orders.xlsx").write_bytes(template("jobs"))
    url = f"sqlite:///{(tmp_path / 'prod.db').as_posix()}"
    monkeypatch.setenv("FREIGHT_ADMIN_PASSWORD", "a-strong-first-pass")
    code = main(
        [
            "create-company",
            "--name",
            "Gulf Haulage",
            "--admin-email",
            "Ops@Gulf.sa",
            "--admin-name",
            "Ops lead",
            "--fleet",
            str(tmp_path / "fleet.xlsx"),
            "--orders",
            str(tmp_path / "orders.xlsx"),
            "--start",
            "2026-10-19 06:00",
            "--database",
            url,
        ]
    )
    assert code == 0
    app = create_app(url, ThreadPoolExecutor(1))
    with TestClient(app) as client:
        signed_in = client.post(
            "/api/auth/login",
            json={"email": "ops@gulf.sa", "password": "a-strong-first-pass"},
            headers=WRITE,
        )
        assert signed_in.json()["user"]["role"] == "admin"
        data = client.get("/api/data").json()
        assert data["fleet"]["trucks"] == 2 and data["jobs"]["contract"] + data["jobs"]["spot"] == 2
        assert data["window"]["start"].startswith("2026-10-19T03:00")  # 06:00 Riyadh
        sheet = load_workbook(io.BytesIO(client.get("/api/data/fleet.xlsx").content))
        assert sheet["Fleet"].max_row == 3


def test_a_bad_fleet_file_creates_nothing(tmp_path, monkeypatch):
    (tmp_path / "fleet.csv").write_text("truck_id,home_site,trailer_type\nX-1,Atlantis,dry_van\n")
    url = f"sqlite:///{(tmp_path / 'prod.db').as_posix()}"
    monkeypatch.setenv("FREIGHT_ADMIN_PASSWORD", "a-strong-first-pass")
    code = main(
        [
            "create-company",
            "--name",
            "Broken",
            "--admin-email",
            "a@b.sa",
            "--admin-name",
            "A",
            "--fleet",
            str(tmp_path / "fleet.csv"),
            "--start",
            "2026-10-19 06:00",
            "--database",
            url,
        ]
    )
    assert code == 1
    assert not (tmp_path / "prod.db").exists() or _companies(url) == 0


def _companies(url: str) -> int:
    from freight.api.db import Company, Database

    with Database(url).session() as session:
        return session.query(Company).count()


def test_an_unreadable_start_time_is_refused(tmp_path, monkeypatch):
    monkeypatch.setenv("FREIGHT_ADMIN_PASSWORD", "a-strong-first-pass")
    url = f"sqlite:///{(tmp_path / 'x.db').as_posix()}"
    args = ["--admin-email", "a@b.sa", "--admin-name", "A", "--fleet", "missing.xlsx"]
    assert (
        main(["create-company", "--name", "X", *args, "--start", "tomorrow", "--database", url])
        == 2
    )
