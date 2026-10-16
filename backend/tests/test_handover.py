import io

from openpyxl import load_workbook

from freight.domain.models import AssignmentLock
from freight.fixtures import demo_snapshot
from freight.handover import export_workbook, plan_changes
from freight.planning.engine import plan_lns, routes_of


def two_plans():
    snapshot = demo_snapshot(job_count=8, team_count=4)
    before = plan_lns(snapshot, rounds=200, seed=1)
    moved = before.assignments[0]
    target = next(t.id for t in snapshot.teams if t.id != moved.team_id)
    after_snapshot = snapshot.model_copy(
        update={"locks": (AssignmentLock(job_id=moved.job_id, team_id=target),)}
    )
    after = plan_lns(after_snapshot, rounds=200, seed=1, initial=routes_of(before))
    return snapshot, before, after_snapshot, after, moved.job_id


def test_changes_name_the_moved_job_and_the_trucks_to_tell():
    snapshot, before, after_snapshot, after, job_id = two_plans()
    changes = plan_changes(after_snapshot, after, snapshot, before)
    served = {a.job_id for a in after.assignments}
    if job_id in served:
        (moved,) = [m for m in changes["moved"] if m["job"] == job_id]
        trucks = {row["truck"] for row in changes["notify"]}
        assert {moved["from"], moved["to"]} <= trucks
    else:
        assert job_id in {d["job"] for d in changes["dropped"]}
    counted = len(changes["moved"]) + len(changes["retimed"]) + len(changes["added"])
    assert counted + changes["unchanged"] == len(after.assignments)


def test_export_has_every_activity_in_both_languages():
    snapshot, before, after_snapshot, after, _ = two_plans()
    changes = plan_changes(after_snapshot, after, snapshot, before)
    english = load_workbook(io.BytesIO(export_workbook(after_snapshot, after, changes, "en")))
    assert english.sheetnames == ["Summary", "Instructions", "Trucks", "Changes", "Unserved"]
    assert english["Instructions"].max_row - 1 == len(after.activities)
    assert english["Trucks"].max_row - 1 == len(after_snapshot.teams)
    arabic = load_workbook(io.BytesIO(export_workbook(after_snapshot, after, None, "ar")))
    assert arabic.sheetnames == ["الملخص", "التعليمات", "الشاحنات", "غير منفذة"]
    assert arabic["التعليمات"].sheet_view.rightToLeft
