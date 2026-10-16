"""Orders and fleet in and out of a snapshot as Excel or CSV sheets.

A dispatcher downloads the current data, edits it in Excel and uploads it back.
Every problem is reported per cell (row, column, code) before anything changes,
so the web page can show it in Arabic or English.

Times without a timezone are Riyadh time. Sites can be named by id, English
name or Arabic name.
"""

from __future__ import annotations

import csv
import io
import re
from collections.abc import Callable, Iterable
from dataclasses import dataclass, field
from datetime import UTC, date, datetime, timedelta
from typing import Any, Literal
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import ValidationError

from freight.domain.models import Job, Resource, Snapshot, Team, Window

RIYADH = ZoneInfo("Asia/Riyadh")
Kind = Literal["jobs", "fleet"]
MAX_ROWS = {"jobs": 2000, "fleet": 300}
AVAILABILITY_TAIL = timedelta(days=2)  # trucks may finish work after the window ends

JOB_COLUMNS = (
    "reference",
    "customer",
    "type",
    "pickup_site",
    "delivery_site",
    "pickup_from",
    "pickup_to",
    "delivery_from",
    "delivery_to",
    "trailer_type",
    "weight_kg",
    "price_sar",
    "loading_minutes",
    "unloading_minutes",
    "turnaround_minutes",
    "priority",
)
FLEET_COLUMNS = (
    "truck_id",
    "truck_name",
    "driver_name",
    "home_site",
    "current_site",
    "trailer_type",
    "capacity_kg",
    "available_from",
    "driver_time_off",
    "truck_workshop",
)
REQUIRED = {
    "jobs": {
        "reference",
        "customer",
        "pickup_site",
        "delivery_site",
        "pickup_from",
        "pickup_to",
        "delivery_from",
        "delivery_to",
        "trailer_type",
        "weight_kg",
    },
    "fleet": {"truck_id", "home_site", "trailer_type"},
}
TRAILERS = {
    "dry_van": "dry_van",
    "dry van": "dry_van",
    "van": "dry_van",
    "مغلقة": "dry_van",
    "flatbed": "flatbed",
    "flat": "flatbed",
    "سطحة": "flatbed",
    "مسطحة": "flatbed",
}
JOB_TYPES = {
    "contract": "confirmed",
    "confirmed": "confirmed",
    "عقد": "confirmed",
    "spot": "offered",
    "offer": "offered",
    "offered": "offered",
    "فوري": "offered",
}
TIME_FORMATS = ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M", "%d/%m/%Y %H:%M")


@dataclass(frozen=True)
class Problem:
    """One problem in an uploaded sheet. ``row`` is the spreadsheet row (header = 1)."""

    row: int
    column: str
    code: str
    value: str = ""
    detail: str = ""


@dataclass
class Outcome:
    snapshot: Snapshot | None
    problems: list[Problem] = field(default_factory=list)
    summary: dict[str, int] = field(default_factory=dict)


class SheetError(ValueError):
    """The file itself cannot be read (wrong type, no header, too many rows)."""

    def __init__(self, code: str, detail: str = "") -> None:
        super().__init__(code)
        self.code = code
        self.detail = detail


# --- reading ---------------------------------------------------------------------------


def _header(value: object) -> str:
    return re.sub(r"[\s-]+", "_", str(value or "").strip().lower())


def read_rows(filename: str, content: bytes, kind: Kind) -> list[tuple[int, dict[str, Any]]]:
    """Rows of a CSV or Excel file as ``(row number, {column: value})``; blank rows skipped."""
    name = filename.lower()
    if name.endswith(".xlsx"):
        try:
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as error:  # openpyxl raises many types for a damaged file
            raise SheetError("unreadable", str(error)) from error
        sheet = book.worksheets[0]
        table: list[list[Any]] = [list(row) for row in sheet.iter_rows(values_only=True)]
    elif name.endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise SheetError("encoding") from error
        table = [list(row) for row in csv.reader(io.StringIO(text))]
    else:
        raise SheetError("file_type")
    if not table:
        raise SheetError("no_header")
    header = [_header(cell) for cell in table[0]]
    known = set(JOB_COLUMNS if kind == "jobs" else FLEET_COLUMNS)
    missing = REQUIRED[kind] - set(header)
    if missing:
        raise SheetError("missing_columns", ", ".join(sorted(missing)))
    rows = []
    for number, cells in enumerate(table[1:], start=2):
        values = {
            column: cell
            for column, cell in zip(header, cells, strict=False)
            if column in known and cell is not None and str(cell).strip() != ""
        }
        if values:
            rows.append((number, values))
    if len(rows) > MAX_ROWS[kind]:
        raise SheetError("too_many_rows", str(MAX_ROWS[kind]))
    return rows


class _Row:
    """Typed access to one row's cells, collecting problems instead of raising."""

    def __init__(self, line: int, values: dict[str, Any], problems: list[Problem]) -> None:
        self.line, self.values, self.problems = line, values, problems
        self.ok = True

    def problem(self, column: str, code: str, value: object = "", detail: str = "") -> None:
        self.ok = False
        self.problems.append(Problem(self.line, column, code, str(value), detail))

    def text(self, column: str, *, required: bool = False) -> str | None:
        value = self.values.get(column)
        if value is None:
            if required:
                self.problem(column, "required")
            return None
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        return str(value).strip()

    def number(
        self, column: str, *, low: float, high: float, required: bool = False
    ) -> float | None:
        raw = self.values.get(column)
        if raw is None:
            if required:
                self.problem(column, "required")
            return None
        try:
            value = float(str(raw).replace(",", "").strip())
        except ValueError:
            self.problem(column, "bad_number", raw)
            return None
        if not low <= value <= high:
            self.problem(column, "out_of_range", raw, f"{low:g}–{high:g}")
            return None
        return value

    def choice(
        self, column: str, options: dict[str, str], *, default: str | None = None
    ) -> str | None:
        raw = self.text(column, required=default is None)
        if raw is None:
            return default
        found = options.get(raw.lower()) or options.get(raw)
        if found is None:
            self.problem(column, "bad_choice", raw, ", ".join(sorted(set(options.values()))))
        return found

    def time(self, column: str, *, required: bool = False) -> datetime | None:
        raw = self.values.get(column)
        if raw is None:
            if required:
                self.problem(column, "required")
            return None
        parsed = parse_time(raw)
        if parsed is None:
            self.problem(column, "bad_time", raw)
        return parsed


def _read_time(raw: object) -> datetime | None:
    if isinstance(raw, datetime):
        return raw
    if isinstance(raw, date):
        return datetime(raw.year, raw.month, raw.day)
    text = str(raw).strip()
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        pass
    for pattern in TIME_FORMATS:
        try:
            return datetime.strptime(text, pattern)
        except ValueError:
            continue
    return None


def parse_time(raw: object) -> datetime | None:
    """A cell as an aware UTC time; naive values are Riyadh time."""
    value = _read_time(raw)
    if value is None:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=RIYADH)
    return value.astimezone(UTC)


def _windows(raw: str) -> list[Window] | None:
    """``2026-10-20`` (whole Riyadh day) or ``2026-10-19 18:00 to 2026-10-20 18:00``, ``;``-separated."""
    result = []
    for part in filter(None, (p.strip() for p in raw.split(";"))):
        if " to " in part:
            start, end = (parse_time(p) for p in part.split(" to ", 1))
        else:
            start = parse_time(part) if re.fullmatch(r"\d{4}-\d{2}-\d{2}", part) else None
            end = start + timedelta(days=1) if start else None
        if start is None or end is None or end <= start:
            return None
        result.append(Window(start=start, end=end))
    return result


def _time_off(row: _Row, column: str) -> tuple[Window, ...]:
    text = row.text(column)
    if not text:
        return ()
    windows = _windows(text)
    if windows is None:
        row.problem(column, "bad_time", text)
        return ()
    return tuple(windows)


def _site_lookup(snapshot: Snapshot) -> Callable[[str], str | None]:
    names: dict[str, str] = {}
    for site in snapshot.sites:
        for key in (site.id, site.name, site.name_ar):
            if key:
                names[key.strip().lower()] = site.id
    return lambda value: names.get(value.strip().lower())


def _slug(value: str) -> str:
    return re.sub(r"[^\w.-]+", "-", value).strip("-")[:90] or "x"


# --- importing -------------------------------------------------------------------------


def import_jobs(snapshot: Snapshot, rows: Iterable[tuple[int, dict[str, Any]]]) -> Outcome:
    """Replace the snapshot's orders with the uploaded ones (all or nothing)."""
    problems: list[Problem] = []
    site = _site_lookup(snapshot)
    existing = {job.reference: job.id for job in snapshot.jobs}
    jobs: list[Job] = []
    seen: dict[str, int] = {}
    used_ids: set[str] = set()
    for number, values in rows:
        row = _Row(number, values, problems)
        reference = row.text("reference", required=True)
        if reference is not None:
            if reference in seen:
                row.problem("reference", "duplicate", reference, str(seen[reference]))
            seen[reference] = number
        customer = row.text("customer", required=True)
        state = row.choice("type", JOB_TYPES, default="confirmed")
        sites = {}
        for column in ("pickup_site", "delivery_site"):
            name = row.text(column, required=True)
            if name is not None:
                sites[column] = site(name)
                if sites[column] is None:
                    row.problem(column, "unknown_site", name)
        times = {
            column: row.time(column, required=True)
            for column in ("pickup_from", "pickup_to", "delivery_from", "delivery_to")
        }
        trailer = row.choice("trailer_type", TRAILERS)
        weight = row.number("weight_kg", low=1, high=100_000, required=True)
        price = row.number("price_sar", low=0, high=10_000_000)
        handling = {
            column: row.number(column, low=0, high=480)
            for column in ("loading_minutes", "unloading_minutes", "turnaround_minutes")
        }
        priority = row.number("priority", low=1, high=3)
        for start, end in (("pickup_from", "pickup_to"), ("delivery_from", "delivery_to")):
            first, last = times[start], times[end]
            if first and last and last < first:
                row.problem(end, "ends_before_start", values.get(end, ""))
        first, last = times["pickup_from"], times["delivery_to"]
        if first and last and last < first:
            row.problem("delivery_to", "delivery_before_pickup", values.get("delivery_to", ""))
        if not row.ok or reference is None:
            continue
        job_id = existing.get(reference) or f"job-{_slug(reference)}"
        while job_id in used_ids:
            job_id += "-2"
        used_ids.add(job_id)
        defaults = {"loading_minutes": 30, "unloading_minutes": 30, "turnaround_minutes": 15}
        try:
            job = Job.model_validate(
                {
                    "id": job_id,
                    "reference": reference,
                    "customer": customer,
                    "pickup_site_id": sites["pickup_site"],
                    "delivery_site_id": sites["delivery_site"],
                    "pickup_window": {"start": times["pickup_from"], "end": times["pickup_to"]},
                    "delivery_window": {
                        "start": times["delivery_from"],
                        "end": times["delivery_to"],
                    },
                    "equipment": trailer,
                    "weight_kg": round(weight or 0),
                    "state": state,
                    "revenue_minor": None if price is None else round(price * 100),
                    "priority": round(priority or 1),
                }
                | {k: defaults[k] if v is None else round(v) for k, v in handling.items()}
            )
        except ValidationError as error:
            row.problem("", "invalid", "", error.errors()[0]["msg"])
            continue
        jobs.append(job)
    summary = {
        "rows": len(seen),
        "contract": sum(j.state == "confirmed" for j in jobs),
        "spot": sum(j.state == "offered" for j in jobs),
        "kept": sum(1 for j in jobs if j.reference in existing),
        "added": sum(1 for j in jobs if j.reference not in existing),
        "removed": len(set(existing) - set(seen)),
        # Later orders wait for a later plan; expired ones can no longer be picked up.
        "later": sum(j.pickup_window.start >= snapshot.horizon.end for j in jobs),
        "expired": sum(
            j.state == "confirmed" and j.pickup_window.end <= snapshot.horizon.start for j in jobs
        ),
    }
    if problems:
        return Outcome(None, problems, summary)
    return _rebuild(snapshot, {"jobs": tuple(jobs)}, summary)


def import_fleet(snapshot: Snapshot, rows: Iterable[tuple[int, dict[str, Any]]]) -> Outcome:
    """Replace the snapshot's trucks, trailers, drivers and teams with the uploaded fleet."""
    problems: list[Problem] = []
    site = _site_lookup(snapshot)
    resources: list[Resource] = []
    teams: list[Team] = []
    seen: dict[str, int] = {}
    before = {team.id for team in snapshot.teams}
    start, end = snapshot.horizon.start, snapshot.horizon.end + AVAILABILITY_TAIL
    for number, values in rows:
        row = _Row(number, values, problems)
        truck = row.text("truck_id", required=True)
        if truck is not None:
            if not re.fullmatch(r"[\w.-]{1,60}", truck):
                row.problem("truck_id", "bad_id", truck)
            elif truck in seen:
                row.problem("truck_id", "duplicate", truck, str(seen[truck]))
            seen[truck] = number
        home_name = row.text("home_site", required=True)
        home = site(home_name) if home_name else None
        if home_name and home is None:
            row.problem("home_site", "unknown_site", home_name)
        current_name = row.text("current_site")
        current = site(current_name) if current_name else home
        if current_name and current is None:
            row.problem("current_site", "unknown_site", current_name)
        trailer = row.choice("trailer_type", TRAILERS)
        capacity = row.number("capacity_kg", low=1, high=100_000)
        available = row.time("available_from")
        if available and available >= snapshot.horizon.end:
            row.problem("available_from", "after_window", values.get("available_from", ""))
        time_off = _time_off(row, "driver_time_off")
        workshop = _time_off(row, "truck_workshop")
        if not row.ok or truck is None or home is None or current is None:
            continue
        window = Window(start=max(start, available) if available else start, end=end)
        stamp = snapshot.horizon.start
        common = {"site_id": current, "availability": window, "confirmed_at": stamp}
        resources += [
            Resource.model_validate(
                common
                | {
                    "id": f"truck-{truck}",
                    "kind": "tractor",
                    "unavailable": workshop,
                    "name": row.text("truck_name") or f"Truck {truck}",
                }
            ),
            Resource.model_validate(
                common
                | {
                    "id": f"trailer-{truck}",
                    "kind": "trailer",
                    "name": f"Trailer {truck}",
                    "equipment": trailer,
                    "capacity_kg": round(capacity or 25_000),
                }
            ),
            Resource.model_validate(
                common
                | {
                    "id": f"driver-{truck}",
                    "kind": "driver",
                    "name": row.text("driver_name") or f"Driver {truck}",
                    "unavailable": time_off,
                }
            ),
        ]
        teams.append(
            Team(
                id=f"team-{truck}",
                name=f"Team {truck}",
                tractor_id=f"truck-{truck}",
                trailer_id=f"trailer-{truck}",
                driver_id=f"driver-{truck}",
                home_site_id=home if current != home else None,
            )
        )
    after = {team.id for team in teams}
    summary = {
        "rows": len(seen),
        "trucks": len(teams),
        "kept": len(after & before),
        "added": len(after - before),
        "removed": len(before - after),
    }
    if problems:
        return Outcome(None, problems, summary)
    return _rebuild(snapshot, {"resources": tuple(resources), "teams": tuple(teams)}, summary)


def _rebuild(snapshot: Snapshot, changes: dict[str, Any], summary: dict[str, int]) -> Outcome:
    """Apply the new rows and drop locks, rules and history that point at removed ids."""
    data = dict(snapshot) | changes  # field -> value, keeping the model objects
    jobs = {job.id for job in data["jobs"]}
    teams = {team.id for team in data["teams"]}
    resources = {resource.id for resource in data["resources"]}
    customers = {job.customer for job in data["jobs"]}
    locks = [lock for lock in snapshot.locks if lock.job_id in jobs and lock.team_id in teams]
    rules = [
        rule
        for rule in snapshot.rules
        if (rule.job_id is None or rule.job_id in jobs)
        and (rule.team_id is None or rule.team_id in teams)
        and (rule.resource_id is None or rule.resource_id in resources)
        and (rule.customer is None or rule.customer in customers)
    ]
    data["locks"] = locks
    data["rules"] = rules
    data["previous"] = {j: t for j, t in snapshot.previous.items() if j in jobs and t in teams}
    summary = summary | {
        "locks_dropped": len(snapshot.locks) - len(locks),
        "rules_dropped": len(snapshot.rules) - len(rules),
    }
    try:
        return Outcome(Snapshot.model_validate(data), [], summary)
    except ValidationError as error:
        problem = Problem(0, "", "invalid", "", error.errors()[0]["msg"])
        return Outcome(None, [problem], summary)


def set_window(snapshot: Snapshot, start: datetime, days: int) -> Snapshot:
    """Move the planning window; trucks available "from the start" move with it."""
    horizon = Window(start=start, end=start + timedelta(days=days))
    old = snapshot.horizon.start
    resources = []
    for resource in snapshot.resources:
        begin = resource.availability.start
        begin = horizon.start if begin <= old else max(begin, horizon.start)
        end = max(begin, horizon.end + AVAILABILITY_TAIL)
        resources.append(
            resource.model_copy(
                update={"availability": Window(start=begin, end=end), "confirmed_at": horizon.start}
            )
        )
    data = dict(snapshot) | {"horizon": horizon, "resources": tuple(resources)}
    return Snapshot.model_validate(data)


# --- exporting -------------------------------------------------------------------------


def _local(value: datetime) -> datetime:
    """Riyadh wall-clock time without tzinfo (Excel cannot store timezones)."""
    return value.astimezone(RIYADH).replace(tzinfo=None)


def job_rows(snapshot: Snapshot) -> list[list[Any]]:
    sites = {site.id: site.name for site in snapshot.sites}
    return [
        [
            job.reference,
            job.customer,
            "contract" if job.state == "confirmed" else "spot",
            sites[job.pickup_site_id],
            sites[job.delivery_site_id],
            _local(job.pickup_window.start),
            _local(job.pickup_window.end),
            _local(job.delivery_window.start),
            _local(job.delivery_window.end),
            job.equipment,
            job.weight_kg,
            None if job.revenue_minor is None else job.revenue_minor / 100,
            job.loading_minutes,
            job.unloading_minutes,
            job.turnaround_minutes,
            job.priority,
        ]
        for job in snapshot.jobs
        if job.state != "cancelled"
    ]


def fleet_rows(snapshot: Snapshot) -> list[list[Any]]:
    sites = {site.id: site.name for site in snapshot.sites}
    resources = {resource.id: resource for resource in snapshot.resources}
    rows = []
    for team in snapshot.teams:
        tractor, trailer, driver = (resources[r] for r in team.resource_ids)
        truck = tractor.id.removeprefix("truck-")

        def spans(windows: tuple[Window, ...]) -> str | None:
            text = "; ".join(
                f"{_local(w.start):%Y-%m-%d %H:%M} to {_local(w.end):%Y-%m-%d %H:%M}"
                for w in windows
            )
            return text or None

        available = tractor.availability.start
        rows.append(
            [
                truck,
                tractor.name,
                driver.name,
                sites[snapshot.home_of(team)],
                sites[tractor.site_id] if tractor.site_id != snapshot.home_of(team) else None,
                trailer.equipment,
                trailer.capacity_kg,
                _local(available) if available > snapshot.horizon.start else None,
                spans(driver.unavailable),
                spans(tractor.unavailable),
            ]
        )
    return rows


EXAMPLES: dict[str, list[list[Any]]] = {
    "jobs": [
        [
            "SO-1001",
            "Customer A",
            "contract",
            "Riyadh 2nd Industrial City",
            "Jeddah Islamic Port",
            "2026-10-19 08:00",
            "2026-10-19 14:00",
            "2026-10-20 08:00",
            "2026-10-20 20:00",
            "dry_van",
            22000,
            4200,
            60,
            60,
            30,
            1,
        ],
        [
            "SPOT-77",
            "Broker B",
            "spot",
            "Jeddah Industrial City",
            "الرياض",
            "2026-10-20 06:00",
            "2026-10-20 18:00",
            "2026-10-21 06:00",
            "2026-10-21 22:00",
            "flatbed",
            18000,
            3100,
            None,
            None,
            None,
            None,
        ],
    ],
    "fleet": [
        [
            "RUH-001",
            "Volvo FH 4521",
            "Ahmed",
            "Riyadh 2nd Industrial City",
            None,
            "dry_van",
            25000,
            None,
            "2026-10-23",
            None,
        ],
        [
            "JED-004",
            None,
            None,
            "Jeddah Industrial City",
            "Dammam 2nd Industrial City",
            "flatbed",
            30000,
            "2026-10-19 12:00",
            "2026-10-21 18:00 to 2026-10-22 18:00",
            "2026-10-24",
        ],
    ],
}

HELP = {
    "jobs": [
        (
            "reference",
            "Your order number (unique). Re-uploading the same number keeps its plan history.",
            "رقم الطلب لديكم (فريد).",
        ),
        ("customer", "Customer name.", "اسم العميل."),
        (
            "type",
            "contract (must be served) or spot (take only if it pays). Default: contract.",
            "contract للشحنات المتعاقد عليها أو spot للحمولات الفورية.",
        ),
        (
            "pickup_site / delivery_site",
            "Site id, English or Arabic name; see the Sites sheet.",
            "رمز الموقع أو اسمه بالعربية أو الإنجليزية؛ راجع ورقة Sites.",
        ),
        (
            "pickup_from … delivery_to",
            "Riyadh time, e.g. 2026-10-19 08:00 or 19/10/2026 08:00.",
            "بتوقيت الرياض، مثل 2026-10-19 08:00.",
        ),
        ("trailer_type", "dry_van or flatbed.", "dry_van (مغلقة) أو flatbed (سطحة)."),
        ("weight_kg", "Load weight in kilograms.", "وزن الحمولة بالكيلوغرام."),
        (
            "price_sar",
            "What the customer pays (optional; needed to value spot loads).",
            "المبلغ الذي يدفعه العميل بالريال (اختياري).",
        ),
        (
            "loading/unloading/turnaround_minutes",
            "Optional. Defaults 30 / 30 / 15.",
            "اختياري. الافتراضي 30 / 30 / 15 دقيقة.",
        ),
        (
            "priority",
            "1 (normal) to 3 (most important). Optional.",
            "من 1 (عادي) إلى 3 (الأهم). اختياري.",
        ),
    ],
    "fleet": [
        (
            "truck_id",
            "Your truck number (letters, digits, - . _). One row per truck with its driver and trailer.",
            "رقم الشاحنة لديكم. صف واحد لكل شاحنة مع سائقها ومقطورتها.",
        ),
        ("truck_name / driver_name", "Optional display names.", "أسماء للعرض (اختياري)."),
        ("home_site", "The depot the truck belongs to.", "المقر الذي تتبع له الشاحنة."),
        (
            "current_site",
            "Where the truck is now, if not at home. Optional.",
            "مكان الشاحنة الآن إن لم تكن في مقرها (اختياري).",
        ),
        ("trailer_type", "dry_van or flatbed.", "dry_van (مغلقة) أو flatbed (سطحة)."),
        (
            "capacity_kg",
            "Trailer capacity. Default 25000.",
            "حمولة المقطورة القصوى. الافتراضي 25000.",
        ),
        (
            "available_from",
            "When the truck can start, if later than the planning window. Optional.",
            "متى تبدأ الشاحنة العمل إن كان بعد بداية نافذة التخطيط (اختياري).",
        ),
        (
            "truck_workshop",
            "Days the truck is in the workshop; same format as driver_time_off.",
            "أيام وجود الشاحنة في الورشة، بنفس صيغة إجازات السائق.",
        ),
        (
            "driver_time_off",
            "Days off: 2026-10-23 (whole day) or 2026-10-21 18:00 to 2026-10-22 18:00; separate with ;",
            "إجازات السائق: تاريخ كامل أو فترة من … to …؛ افصل بينها بـ ;",
        ),
    ],
}


def workbook(snapshot: Snapshot, kind: Kind, *, template: bool = False) -> bytes:
    """An Excel file with the data (or examples), a Sites sheet and a bilingual Help sheet."""
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Orders" if kind == "jobs" else "Fleet"
    columns = JOB_COLUMNS if kind == "jobs" else FLEET_COLUMNS
    sheet.append(list(columns))
    rows = (
        EXAMPLES[kind]
        if template
        else job_rows(snapshot)
        if kind == "jobs"
        else fleet_rows(snapshot)
    )
    for row in rows:
        sheet.append(row)
    bold, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="176858")
    for cell in sheet[1]:
        cell.font, cell.fill = bold, fill
    for index, column in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(14, len(column) + 4)
        if column.endswith(("_from", "_to")):
            for line in range(2, sheet.max_row + 1):
                sheet.cell(row=line, column=index).number_format = "yyyy-mm-dd hh:mm"
    sheet.freeze_panes = "A2"

    sites = book.create_sheet("Sites")
    sites.append(["id", "name", "name_ar"])
    for site in snapshot.sites:
        sites.append([site.id, site.name, site.name_ar])
    for cell in sites[1]:
        cell.font, cell.fill = bold, fill
    for letter, width in (("A", 18), ("B", 34), ("C", 40)):
        sites.column_dimensions[letter].width = width

    help_sheet = book.create_sheet("Help")
    help_sheet.append(["column", "meaning", "المعنى"])
    for entry in HELP[kind]:
        help_sheet.append(list(entry))
    for cell in help_sheet[1]:
        cell.font, cell.fill = bold, fill
    for letter, width in (("A", 30), ("B", 80), ("C", 70)):
        help_sheet.column_dimensions[letter].width = width

    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
