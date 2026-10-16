"""Handing a plan over: what changed since the last plan, and an Excel file to send.

The changes are in the terms a dispatcher acts on: which jobs moved to another
truck, which pickups moved in time, what is newly planned or no longer planned,
and so which trucks need new instructions.
"""

from __future__ import annotations

import io
from collections import defaultdict
from datetime import datetime
from typing import Any, Literal
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from freight.domain.models import Activity, Plan, Snapshot
from freight.planning.diff import diff_plans

RIYADH = ZoneInfo("Asia/Riyadh")
Language = Literal["en", "ar"]


def plan_changes(
    snapshot: Snapshot, plan: Plan, before_snapshot: Snapshot, before: Plan
) -> dict[str, Any]:
    """Changes from ``before`` to ``plan``, with names and references for display."""
    diff = diff_plans(before, plan)
    teams = {t.id: t.name for t in (*before_snapshot.teams, *snapshot.teams)}
    jobs = {j.id: j for j in (*before_snapshot.jobs, *snapshot.jobs)}
    still_open = {j.id for j in snapshot.jobs if j.state != "cancelled"}

    def job(job_id: str) -> dict[str, Any]:
        found = jobs.get(job_id)
        return {
            "job": job_id,
            "ref": found.reference if found else job_id,
            "customer": found.customer if found else "",
        }

    moved = [job(j) | {"from": teams.get(a, a), "to": teams.get(b, b)} for j, a, b in diff.moved]
    retimed = [
        job(j) | {"truck": teams.get(t, t), "before": x.isoformat(), "after": y.isoformat()}
        for j, t, x, y in diff.retimed
    ]
    added = [job(j) | {"truck": teams.get(t, t)} for j, t in diff.added]
    # A job that left the order book is not a change to act on; an open job that lost its
    # truck is.
    dropped = [job(j) | {"truck": teams.get(t, t)} for j, t in diff.dropped if j in still_open]
    touched: defaultdict[str, int] = defaultdict(int)
    for _, a, b in diff.moved:
        touched[a] += 1
        touched[b] += 1
    for _, t, _, _ in diff.retimed:
        touched[t] += 1
    for j, t in diff.dropped:
        if j in still_open:
            touched[t] += 1
    return {
        "moved": moved,
        "retimed": retimed,
        "added": added,
        "dropped": dropped,
        "unchanged": diff.unchanged,
        # Trucks whose already-planned work changed: their drivers must be told. New work
        # on top (orders that entered the window) is ordinary dispatch, counted apart.
        "notify": sorted(
            ({"truck": teams.get(t, t), "changes": n} for t, n in touched.items()),
            key=lambda row: str(row["truck"]),
        ),
        "new_work_trucks": len({t for _, t in diff.added} - set(touched)),
    }


# --- Excel export ----------------------------------------------------------------------

WORDS: dict[str, dict[Language, str]] = {
    "summary": {"en": "Summary", "ar": "الملخص"},
    "instructions": {"en": "Instructions", "ar": "التعليمات"},
    "trucks": {"en": "Trucks", "ar": "الشاحنات"},
    "changes": {"en": "Changes", "ar": "التغييرات"},
    "unserved": {"en": "Unserved", "ar": "غير منفذة"},
    "truck": {"en": "Truck", "ar": "الشاحنة"},
    "driver": {"en": "Driver", "ar": "السائق"},
    "depot": {"en": "Depot", "ar": "المقر"},
    "start": {"en": "Start (Riyadh)", "ar": "البداية (الرياض)"},
    "end": {"en": "End (Riyadh)", "ar": "النهاية (الرياض)"},
    "activity": {"en": "Activity", "ar": "النشاط"},
    "from": {"en": "From", "ar": "من"},
    "to": {"en": "To", "ar": "إلى"},
    "km": {"en": "km", "ar": "كم"},
    "job": {"en": "Order", "ar": "الطلب"},
    "customer": {"en": "Customer", "ar": "العميل"},
    "breaks": {"en": "Breaks / rests", "ar": "استراحات / راحات"},
    "jobs": {"en": "Orders", "ar": "الطلبات"},
    "loaded_km": {"en": "Loaded km", "ar": "كم محمّلة"},
    "empty_km": {"en": "Empty km", "ar": "كم فارغة"},
    "revenue": {"en": "Revenue (SAR)", "ar": "الإيراد (ريال)"},
    "driving_h": {"en": "Driving hours", "ar": "ساعات القيادة"},
    "change": {"en": "Change", "ar": "التغيير"},
    "detail": {"en": "Detail", "ar": "التفاصيل"},
    "moved": {"en": "Moved to another truck", "ar": "نُقلت إلى شاحنة أخرى"},
    "retimed": {"en": "Pickup re-timed", "ar": "تغيّر وقت التحميل"},
    "added": {"en": "Newly planned", "ar": "مخططة حديثاً"},
    "dropped": {"en": "No longer planned", "ar": "لم تعد مخططة"},
    "pickup_window": {"en": "Pickup window", "ar": "نافذة التحميل"},
    "type": {"en": "Type", "ar": "النوع"},
    "contract": {"en": "contract", "ar": "متعاقد عليها"},
    "spot": {"en": "spot", "ar": "فورية"},
    "item": {"en": "Item", "ar": "البند"},
    "value": {"en": "Value", "ar": "القيمة"},
    "plan_name": {"en": "Operation", "ar": "العملية"},
    "window": {"en": "Planning window", "ar": "نافذة التخطيط"},
    "served": {"en": "Contract orders served", "ar": "الطلبات المتعاقد عليها المنفذة"},
    "offers": {"en": "Spot loads taken", "ar": "الحمولات الفورية المقبولة"},
    "empty_share": {"en": "Share of km driven empty", "ar": "نسبة الكيلومترات الفارغة"},
    "cost": {"en": "Modeled cost (SAR)", "ar": "التكلفة التقديرية (ريال)"},
    "contribution": {"en": "Contribution (SAR)", "ar": "هامش المساهمة (ريال)"},
    "checks": {"en": "Independent checks", "ar": "الفحوصات المستقلة"},
    "passed": {"en": "passed", "ar": "ناجحة"},
    "failed": {"en": "FAILED", "ar": "فاشلة"},
    "loaded": {"en": "Loaded drive", "ar": "قيادة محمّلة"},
    "empty": {"en": "Empty drive", "ar": "قيادة فارغة"},
    "loading": {"en": "Loading", "ar": "تحميل"},
    "unloading": {"en": "Unloading", "ar": "تفريغ"},
    "turnaround": {"en": "Turnaround", "ar": "تجهيز"},
    "waiting": {"en": "Waiting", "ar": "انتظار"},
    "rest": {"en": "Rest", "ar": "راحة"},
}


def _local(value: datetime) -> datetime:
    return value.astimezone(RIYADH).replace(tzinfo=None)


def _sheet(book: Workbook, title: str, header: list[str], rtl: bool) -> Worksheet:
    sheet: Worksheet = book.create_sheet(title)
    sheet.append(header)
    bold, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="176858")
    for cell in sheet[1]:
        cell.font, cell.fill = bold, fill
    sheet.freeze_panes = "A2"
    sheet.sheet_view.rightToLeft = rtl
    return sheet


def _widths(sheet: Worksheet, widths: list[int], times: tuple[int, ...] = ()) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for index in times:
        for line in range(2, sheet.max_row + 1):
            sheet.cell(row=line, column=index).number_format = "yyyy-mm-dd hh:mm"


def export_workbook(
    snapshot: Snapshot,
    plan: Plan,
    changes: dict[str, Any] | None = None,
    language: Language = "en",
) -> bytes:
    """The plan as an Excel file: summary, per-truck instructions, trucks, changes, unserved."""
    w = {key: words[language] for key, words in WORDS.items()}
    rtl = language == "ar"
    sites = {s.id: (s.name_ar or s.name) if rtl else s.name for s in snapshot.sites}
    resources = {r.id: r for r in snapshot.resources}
    jobs = {j.id: j for j in snapshot.jobs}
    teams = sorted(snapshot.teams, key=lambda t: t.name)
    by_team: defaultdict[str, list[Activity]] = defaultdict(list)
    for activity in plan.activities:
        by_team[activity.team_id].append(activity)
    assigned = {a.job_id for a in plan.assignments}
    confirmed = [j for j in snapshot.jobs if j.state == "confirmed"]
    offered = [j for j in snapshot.jobs if j.state == "offered"]
    c = plan.costs

    book = Workbook()
    default = book.active
    assert default is not None
    book.remove(default)

    summary = _sheet(book, w["summary"], [w["item"], w["value"]], rtl)
    empty, loaded = (c.empty_distance_m, c.loaded_distance_m) if c else (0, 0)
    rows: list[tuple[str, Any]] = [
        (w["plan_name"], snapshot.name),
        (
            w["window"],
            f"{_local(snapshot.horizon.start):%Y-%m-%d %H:%M} – "
            f"{_local(snapshot.horizon.end):%Y-%m-%d %H:%M}",
        ),
        (w["served"], f"{sum(j.id in assigned for j in confirmed)} / {len(confirmed)}"),
        (w["offers"], f"{sum(j.id in assigned for j in offered)} / {len(offered)}"),
        (w["empty_share"], f"{round(100 * empty / (empty + loaded))}%" if empty + loaded else "—"),
        (w["cost"], round(c.total_minor / 100) if c and c.total_minor is not None else "—"),
        (
            w["contribution"],
            round((c.revenue_minor - c.total_minor) / 100)
            if c and c.revenue_minor is not None and c.total_minor is not None
            else "—",
        ),
        (w["checks"], w["passed"] if plan.valid else w["failed"]),
    ]
    for row in rows:
        summary.append(list(row))
    _widths(summary, [34, 48])

    instructions = _sheet(
        book,
        w["instructions"],
        [
            w["truck"],
            w["driver"],
            w["start"],
            w["end"],
            w["activity"],
            w["from"],
            w["to"],
            w["km"],
            w["job"],
            w["customer"],
            w["breaks"],
        ],
        rtl,
    )
    trucks = _sheet(
        book,
        w["trucks"],
        [
            w["truck"],
            w["driver"],
            w["depot"],
            w["jobs"],
            w["loaded_km"],
            w["empty_km"],
            w["revenue"],
            w["driving_h"],
        ],
        rtl,
    )
    for team in teams:
        driver = resources[team.driver_id].name
        activities = sorted(by_team.get(team.id, []), key=lambda a: a.start)
        for a in activities:
            job = jobs.get(a.job_id) if a.job_id else None
            instructions.append(
                [
                    team.name,
                    driver,
                    _local(a.start),
                    _local(a.end),
                    w[a.kind],
                    sites[a.from_site_id],
                    sites[a.to_site_id] if a.to_site_id != a.from_site_id else "",
                    round(a.distance_m / 1000) or None,
                    job.reference if job else "",
                    job.customer if job else "",
                    len(a.pauses) or None,
                ]
            )
        team_jobs = {a.job_id for a in activities if a.job_id}
        driving = sum(
            (a.end - a.start).total_seconds() - a.paused_seconds
            for a in activities
            if a.kind in ("loaded", "empty")
        )
        trucks.append(
            [
                team.name,
                driver,
                sites[snapshot.home_of(team)],
                len(team_jobs),
                round(sum(a.distance_m for a in activities if a.kind == "loaded") / 1000),
                round(sum(a.distance_m for a in activities if a.kind == "empty") / 1000),
                round(sum(jobs[j].revenue_minor or 0 for j in team_jobs) / 100),
                round(driving / 3600, 1),
            ]
        )
    _widths(instructions, [16, 18, 17, 17, 16, 30, 30, 8, 14, 26, 14], times=(3, 4))
    _widths(trucks, [16, 18, 30, 9, 11, 11, 14, 13])

    if changes is not None:
        sheet = _sheet(book, w["changes"], [w["change"], w["job"], w["customer"], w["detail"]], rtl)
        for row in changes["moved"]:
            sheet.append([w["moved"], row["ref"], row["customer"], f"{row['from']} → {row['to']}"])
        for row in changes["retimed"]:
            before = _local(datetime.fromisoformat(row["before"]))
            after = _local(datetime.fromisoformat(row["after"]))
            detail = f"{row['truck']}: {before:%a %H:%M} → {after:%a %H:%M}"
            sheet.append([w["retimed"], row["ref"], row["customer"], detail])
        for row in changes["added"]:
            sheet.append([w["added"], row["ref"], row["customer"], row["truck"]])
        for row in changes["dropped"]:
            sheet.append([w["dropped"], row["ref"], row["customer"], row["truck"]])
        _widths(sheet, [26, 14, 26, 48])

    unserved = _sheet(
        book,
        w["unserved"],
        [w["job"], w["customer"], w["type"], w["from"], w["to"], w["pickup_window"]],
        rtl,
    )
    open_unserved = [jobs[j] for j in plan.unserved_job_ids if jobs[j].state != "cancelled"]
    for job in sorted(open_unserved, key=lambda j: (j.state != "confirmed", j.reference)):
        unserved.append(
            [
                job.reference,
                job.customer,
                w["contract"] if job.state == "confirmed" else w["spot"],
                sites[job.pickup_site_id],
                sites[job.delivery_site_id],
                f"{_local(job.pickup_window.start):%Y-%m-%d %H:%M} – "
                f"{_local(job.pickup_window.end):%H:%M}",
            ]
        )
    _widths(unserved, [14, 26, 12, 30, 30, 26])

    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
